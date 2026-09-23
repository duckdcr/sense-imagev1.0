from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE


OUT = "/Users/wangyuanzi/Documents/code/work2/qbh_detail_sence/sense-imagev1.0/docs/01版本白底窗帘生图SOP.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
MUTED = "5B6573"
RED = "9B1C1C"
GOLD = "7A5A00"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


# Exact macOS font face name.  LibreOffice does not resolve the generic
# ``STHeiti`` family reliably in headless mode, which renders Chinese as tofu.
CN_FONT = "STHeiti Light"


def font(run, size=11, bold=False, color=None, name=CN_FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_para_spacing(p, before=0, after=6, line=300):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line / 240


def add_text(doc, text, style=None, bold_prefix=None, color=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        font(r, bold=True, color=color)
        r = p.add_run(text[len(bold_prefix):])
        font(r, color=color)
    else:
        r = p.add_run(text)
        font(r, color=color)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    set_para_spacing(p, after=4, line=300)
    r = p.add_run(text)
    font(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    set_para_spacing(p, after=4, line=300)
    r = p.add_run(text)
    font(r)
    return p


def add_callout(doc, title, text, kind="info"):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [1600, 7760])
    left, right = table.rows[0].cells
    fills = {"info": LIGHT_BLUE, "warn": "FFF4CC", "risk": "FDECEC", "success": "EAF5EF"}
    colors = {"info": DARK_BLUE, "warn": GOLD, "risk": RED, "success": "216E39"}
    set_cell_shading(left, fills[kind])
    set_cell_shading(right, fills[kind])
    p = left.paragraphs[0]
    set_para_spacing(p, after=0, line=280)
    r = p.add_run(title)
    font(r, size=10, bold=True, color=colors[kind])
    p = right.paragraphs[0]
    set_para_spacing(p, after=0, line=280)
    r = p.add_run(text)
    font(r, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths)
    header = table.rows[0]
    set_repeat_header(header)
    for cell, value in zip(header.cells, headers):
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        set_para_spacing(p, after=0, line=280)
        r = p.add_run(value)
        font(r, size=font_size, bold=True, color=INK)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            p = cell.paragraphs[0]
            set_para_spacing(p, after=0, line=280)
            r = p.add_run(value)
            font(r, size=font_size)
    return table


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    if level == 1:
        font(r, size=16, bold=True, color=BLUE)
        set_para_spacing(p, before=18, after=10, line=300)
    elif level == 2:
        font(r, size=13, bold=True, color=BLUE)
        set_para_spacing(p, before=14, after=7, line=300)
    else:
        font(r, size=12, bold=True, color=DARK_BLUE)
        set_para_spacing(p, before=10, after=5, line=300)
    return p


def add_page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = CN_FONT
normal._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
    style = styles[style_name]
    style.font.name = CN_FONT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

# Header/footer - compact_reference_guide.
header_p = section.header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
header_p.paragraph_format.space_after = Pt(0)
r = header_p.add_run("01 版本白底窗帘成品图 SOP")
font(r, size=8, bold=True, color=MUTED)
footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer_p.add_run("内部操作手册  |  2026-08-19")
font(r, size=8, color=MUTED)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(48)
p.paragraph_format.space_after = Pt(8)
r = p.add_run("01 版本")
font(r, size=16, bold=True, color=BLUE)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run("白底窗帘成品图生成 SOP")
font(r, size=26, bold=True, color=INK)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(20)
r = p.add_run("适用于实拍面料图、CSV 花位数据与完整花位图的批量生成流程")
font(r, size=12, color=MUTED)

add_callout(doc, "核心目标", "在纯白背景上输出正视、完整、中间对开的双片窗帘；Image 1 只提供结构，Image 2 只提供颜色和材质，有花位时 Image 3 只提供一个完整花位的二维图形结构。", "info")
add_heading(doc, "适用范围", 1)
add_bullet(doc, "01 版本本地工具：白底图生成页面（默认地址通常为 http://127.0.0.1:3456）。")
add_bullet(doc, "手动上传和数据库批量两种任务方式。")
add_bullet(doc, "32 系列：实拍面料图在 32-10 / 32-15；CSV 管理 SKU、花位纵横尺寸和成分；有花位 SKU 的完整花位图位于居莱/花位图。")
add_heading(doc, "固定成品参数", 1)
add_table(doc, ["项目", "固定值 / 规则"], [
    ("成品闭合尺寸", "整套宽 330 cm × 高 270 cm；左右片成品宽各 165 cm。"),
    ("褶皱", "2 倍褶皱。每片展开布尺寸为 330 cm × 270 cm；两片展开总布宽为 660 cm。"),
    ("中间开口", "约 30 cm；开口区域没有面料，且不额外增加成品额定宽度。"),
    ("左右裁片", "来自同一卷连续面料的相邻连续位置；不可复制、镜像、轴对称或只改亮度。"),
    ("背景", "纯白背景只属于背景；窗帘颜色、花纹与材质只能来自面料输入。")
], [2700, 6660], 9)

add_page_break(doc)
add_heading(doc, "一、生成前的快速判断", 1)
add_callout(doc, "先看 CSV，不凭肉眼猜花位", "CSV 的花位纵向长(cm)和花位横向长(cm)是是否使用完整花位图、以及花回物理尺寸的唯一权威来源。", "warn")
add_heading(doc, "1. 输入资料与职责", 2)
add_table(doc, ["输入", "只负责什么", "绝不负责什么"], [
    ("Image 1\n纯白背景窗帘结构模板", "帘头样式、孔位/捏褶、轮廓、左右布局、开口、褶峰褶谷、垂坠与宏观明暗。", "面料颜色、图案、织纹、材质。"),
    ("Image 2\n实拍面料图", "固有颜色、颜色比例、经纬组织、纱线、颗粒、毛羽、孔隙、粗糙度、光泽、材质密度。", "帘头、开口、褶皱数量和窗帘几何。"),
    ("Image 3\n完整花位图（仅有花位）", "一个完整花位的二维结构、图案方向、元素位置和完整边界。", "颜色、亮度、拍摄光线与材质；这些均以 Image 2 优先。"),
    ("CSV", "SKU、是否有花位、花位纵横厘米尺寸、成分等辅助信息。", "不能替代 Image 2 的真实颜色和织物外观。")
], [1800, 3900, 3660], 8)
add_heading(doc, "2. 两条流程的分流规则", 2)
add_table(doc, ["CSV 花位字段", "自动流程", "提交图片"], [
    ("花位纵向长 = / 且 花位横向长 = /", "无花位 / 素面微纹理", "Image 1 + Image 2。不得发送 Image 3，也不得凭空生成条纹、方格或花位。"),
    ("两个字段均为有效正数", "有花位 / 三图流程", "Image 1 + Image 2 + 对应 SKU 的 Image 3。花位图缺失时必须先补图，不能降级为两图生成。"),
    ("只有一个字段为 /，或字段不是正数", "阻止生成", "先修正 CSV；不得靠提示词猜测花位尺寸。")
], [2400, 2300, 4660], 8)
add_heading(doc, "3. 文件命名和实拍范围", 2)
add_bullet(doc, "文件名含 _detail_10：Image 2 覆盖真实面料 10 cm × 10 cm。")
add_bullet(doc, "文件名含 _detail_15：Image 2 覆盖真实面料 15 cm × 15 cm。")
add_bullet(doc, "实拍范围只锁定纱线、织法、颗粒和单位面积密度；它不是完整花位，也不是可直接平铺的矩形贴图。")
add_bullet(doc, "若文件命名无法识别，手动流程中选择与原图真实覆盖范围一致的 10 cm × 10 cm 或 15 cm × 15 cm。")

add_page_break(doc)
add_heading(doc, "二、无花位 / 素面微纹理 SOP", 1)
add_callout(doc, "适用条件", "CSV 的两个花位字段均为“/”。即使面料看似素色，实拍图里的色纱、经纬组织、孔隙、暗缝和微观反差仍是必须保留的面料身份。", "info")
add_heading(doc, "操作步骤", 2)
steps_plain = [
    "进入“白底图生成”，输入 API Key；选择圆孔帘或韩式双捏褶帘。",
    "上传或选择 Image 2 实拍面料图。确认 SKU、文件名和实拍尺寸（_detail_10 / _detail_15）正确。",
    "预检显示“无花位”后，确认不匹配 Image 3；若 CSV 显示有花位或提示字段异常，停止提交并先核对 CSV。",
    "选择画质和分辨率，提交生成。默认提示词会将 Image 1 的结构与 Image 2 的面料信息分离。",
    "在完成图中核对：颜色没有被白背景洗成米灰；织纹没有被磨平；深色细节没有只出现在褶谷。",
    "合格后移至导出步骤；不合格时重新生成或调整补充提示词，不要把实拍样本当作规则花回手工重复。"
]
for s in steps_plain:
    add_number(doc, s)
add_heading(doc, "无花位验收清单", 2)
for item in [
    "全景颜色仍读取为 Image 2 的综合色相，而不是统一灰、米白、卡其或冷白。",
    "经纬、纱线粗细差、毛羽、孔隙、暗缝、微高光在全景中可自然融合，但不能被磨成平滑纸面或塑料感。",
    "不新增独立花位、横条、竖条、棋盘格、方框或大网格。",
    "褶谷阴影不替代固有深色；褶峰高光不漂白固有颜色。",
    "左右两片同面料、同方向、同密度，但局部纹理不完全镜像。"
]:
    add_bullet(doc, item)

add_page_break(doc)
add_heading(doc, "三、有花位 / 完整花回 SOP", 1)
add_callout(doc, "物理尺度优先", "Image 3 回答“一个花位长什么样”；CSV 回答“这个花位实际多大”；Image 2 回答“这个花位用什么颜色和材质织成”。三者不可互相替代。", "warn")
add_heading(doc, "操作步骤", 2)
steps_repeat = [
    "确认 CSV 的花位纵向长(cm)和花位横向长(cm)均为正数；确认该 SKU 的完整花位图已在居莱/花位图中精确匹配。",
    "上传 Image 2；系统按 SKU 匹配 Image 3。预检必须显示“有花位 X × Y cm · Image 3 已匹配”。",
    "确认成品固定参数：270 cm 高、每片展开宽 330 cm、两片展开总宽 660 cm。",
    "系统在最终提示词中换算：纵向花回数 = 270 ÷ 花位纵向长；每片横向花回数 = 330 ÷ 花位横向长；两片总横向花回数 = 660 ÷ 花位横向长。",
    "提交生成。生成顺序必须是：先建立连续平展花位面料，再裁切成左右两片，最后包覆到 Image 1 的褶皱上。",
    "验收花回数量、方向和连续性；不得仅因花回跨褶皱而放大图案，亦不得在每道褶峰重置花位。"
]
for s in steps_repeat:
    add_number(doc, s)
add_heading(doc, "花回与样本尺度不能混用", 2)
add_table(doc, ["尺度来源", "使用目的", "错误用法"], [
    ("CSV 花位纵横厘米", "决定完整花位在 270 cm / 330 cm / 660 cm 展开布上的数量。", "把 Image 2 里的局部图案数量当作花回数量。"),
    ("Image 3 完整花位图", "决定一个完整花位的轮廓、内部关系、方向与边界。", "把 Image 3 的颜色或拍摄光线覆盖到 Image 2 色号。"),
    ("Image 2 的 10/15 cm 实拍范围", "决定纱线密度、织造颗粒、颜色比例和材质尺度。", "把 10/15 cm 样本拉伸成一整个花回，或当成贴图平铺。")
], [2450, 3610, 3300], 8)
add_heading(doc, "有花位验收清单", 2)
for item in [
    "最终花回数与 CSV 换算一致，允许边缘裁切产生部分花回，但不得明显少循环或大幅缩放。",
    "完整花位跨过褶峰、侧面和褶谷时保持连续，不能在褶谷被切断、重置、镜像或逐褶复制。",
    "Image 3 的结构、Image 2 的颜色与材质同时存在；不能只保留一方。",
    "左右两片来自同一连续卷料的不同位置，花位相位合理但不完全相同。"
]:
    add_bullet(doc, item)

add_page_break(doc)
add_heading(doc, "四、固定尺寸、循环与密度核对", 1)
add_heading(doc, "1. 实拍范围跨度（用于材质密度）", 2)
add_table(doc, ["Image 2 实拍范围", "成品高度 270 cm", "每片展开宽 330 cm", "两片展开总宽 660 cm"], [
    ("10 cm × 10 cm", "270 ÷ 10 = 27 个样本高度跨度", "330 ÷ 10 = 33 个跨度", "660 ÷ 10 = 66 个跨度"),
    ("15 cm × 15 cm", "270 ÷ 15 = 18 个样本高度跨度", "330 ÷ 15 = 22 个跨度", "660 ÷ 15 = 44 个跨度")
], [1900, 2487, 2487, 2486], 8)
add_callout(doc, "重要", "上表只用于保持 Image 2 的材料单位密度：颜色贡献、纱线粗细、孔隙和纹理频率。它不能替代有花位 SKU 的 CSV 花回尺寸。", "risk")
add_heading(doc, "2. 花位数量计算示例", 2)
add_table(doc, ["已知数据", "公式", "结果解释"], [
    ("花位纵向长 V cm", "270 ÷ V", "成品高度内的完整花回数；边缘允许有被裁切的部分花回。"),
    ("花位横向长 H cm", "330 ÷ H", "单片展开布宽度内的横向花回数。"),
    ("花位横向长 H cm", "660 ÷ H", "两片展开总布宽内的横向花回数；对应同一卷连续面料。")
], [2400, 2100, 4860], 8)
add_heading(doc, "3. 快速审图方法", 2)
for item in [
    "先数低频花位或显著横带：若数量明显低于 270 ÷ 纵向周期，说明图案被放大。",
    "再看综合色相：白背景不应把 Image 2 的蓝绿、砖红、黄赭、棕灰等固有颜色洗成统一浅米灰。",
    "最后看材质：微观纱线在远景允许融合，但综合色差、孔隙贡献和粗糙度不应消失。"
]:
    add_bullet(doc, item)

add_page_break(doc)
add_heading(doc, "五、批量生成与导出保存", 1)
add_heading(doc, "1. 手动上传批量", 2)
for s in [
    "在“白底图生成”中选择帘头样式，选择多个实拍面料文件或一个面料文件夹。",
    "确认预检信息：无花位会显示成分（若 CSV 提供）；有花位会显示花位纵向 × 横向厘米和 Image 3 已匹配。",
    "若出现 SKU/花位信息预检失败，先处理 CSV 或完整花位图，再开始生成。",
    "选择画质、分辨率，点击“开始生成”。成功任务会移至“已完成”。"
]:
    add_number(doc, s)
add_heading(doc, "2. 数据库批量", 2)
for s in [
    "切换“数据库批量”，选择任务方式和帘头样式。",
    "先选择保存目录，再点击“读取数据库”；只切换区域不会自动读取数据库。",
    "确认输出格式、并发数和任务数量。网络或上游限流时先降低并发数。",
    "使用“继续未完成”时，工具会根据保存目录中的进度清单跳过已完成 SKU；“重新生成全部”会覆盖同名结果，务必谨慎。"
]:
    add_number(doc, s)
add_heading(doc, "3. 导出已完成结果", 2)
add_callout(doc, "你刚遇到的“保存不了”", "选择批次导出 PNG / WebP 后会先弹出“选择 PNG 批次”或“选择 WebP 批次”。勾选需要导出的批次，再点击底部“选择文件夹并保存 (N)”，随后在系统弹出的文件夹选择器中选定实际目录并确认。未完成最后一步的系统文件夹选择，不会写出文件。", "warn")
for item in [
    "按批次导出：在已完成任务中打开“选择批次导出 PNG/WebP”，勾选批次，点击“选择文件夹并保存”。",
    "单张或已选结果：使用“下载选中”或“导出 WebP”另存。",
    "若看到“输出文件名重复”，先改目标目录或清理同名文件，避免误覆盖。",
    "白底输出命名：圆孔帘为 {sku}_grommet.webp；韩式双捏褶帘为 {sku}_double-pinch.webp。"
]:
    add_bullet(doc, item)

add_page_break(doc)
add_heading(doc, "六、常见问题与处理", 1)
add_table(doc, ["现象", "判断原因", "处理动作"], [
    ("窗帘发灰、发米、颜色被洗淡", "纯白背景或褶皱明暗覆盖了 Image 2 固有色；或历史中使用的是旧提示词。", "确认使用最新默认提示词；检查 Image 2 是否正确；必要时在补充提示词强调保持固有综合色相和颜色面积比例。"),
    ("面料纹理像被磨平", "全景降采样时模型把经纬、孔隙、纱线差异平均掉。", "确认 Image 2 是清晰原图；强调真实织纹、孔隙、毛羽和粗细差自然融合，而非变成平滑素面。"),
    ("有花位图案循环数量不对", "把实拍图范围当作花回，或 CSV 花位尺寸/花位图匹配错误。", "核对 CSV 两个花位厘米字段、对应 SKU 完整花位图和最终换算 270/V、330/H、660/H。"),
    ("有花位却只出现一小段或大幅放大", "Image 3 未匹配或花回尺寸未被正确读取。", "停止生成；检查预检是否显示“Image 3 已匹配”；修复完整花位图文件名或 CSV。"),
    ("无花位却生成条纹/方格", "模型把微织网或样本边界误读为装饰图案。", "确认 CSV 为 / + /；不要发送 Image 3；在补充提示词说明不新增中低频图案。"),
    ("左右两片完全一样", "生成结果出现复制或镜像。", "重新生成；强调左右来自同一卷连续面料的不同连续坐标，禁止镜像/轴对称。"),
    ("专项/预检报错", "CSV 花位字段不完整、实拍尺寸无法识别、或有花位 SKU 缺完整花位图。", "查看错误文字；优先检查 SKU、_detail_10/_detail_15 命名、CSV 和居莱/花位图的精确 SKU 匹配。"),
    ("导出后找不到文件", "只勾选了批次，未在系统文件夹选择器完成保存。", "重新执行批次导出，点击“选择文件夹并保存”，并在系统弹窗中选择实际目录。")
], [2100, 3550, 3710], 8)
add_heading(doc, "提交前最终检查", 2)
for item in [
    "SKU、实拍图、实拍尺寸和帘头样式正确。",
    "无花位：CSV 两个字段均为 /，且只走两图流程。",
    "有花位：CSV 两个字段均为正数，且 Image 3 匹配成功。",
    "颜色、图案、纹理全部以 Image 2 为主；Image 1 只提供结构；Image 3 只提供花位结构。",
    "成品固定尺寸和花回/密度换算均已核对。",
    "导出前确认目标文件夹、格式和重名风险。"
]:
    add_bullet(doc, item)

doc.core_properties.title = "01版本白底窗帘生图SOP"
doc.core_properties.subject = "白底窗帘成品图生成操作标准流程"
doc.core_properties.author = "Codex"
doc.save(OUT)
print(OUT)
