from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches


ROOT = Path("/Users/wangyuanzi/Documents/code/work2/qbh_detail_sence/sense-imagev1.0/docs")
OUT = ROOT / "01版本白底窗帘生图SOP.docx"
PAGE_DIR = ROOT / "_01_sop_pages"

W, H = 1700, 2200
FONT_PATH = "/System/Library/Fonts/Hiragino Sans GB.ttc"
INK = "#102A43"
BLUE = "#1F6FB2"
MUTED = "#536471"
LIGHT = "#EEF4FA"
PALE = "#F7F9FB"
WARN = "#FFF5D6"
RISK = "#FDECEC"
BORDER = "#C8D5E3"
GREEN = "#1E7A46"


def f(size, bold=False):
    return ImageFont.truetype(FONT_PATH, size, index=0)


def wrap(draw, text, font, width):
    lines = []
    for paragraph in str(text).split("\n"):
        current = ""
        for char in paragraph:
            candidate = current + char
            if current and draw.textlength(candidate, font=font) > width:
                lines.append(current)
                current = char
            else:
                current = candidate
        lines.append(current)
    return lines or [""]


class Page:
    def __init__(self, number, title, subtitle=""):
        self.im = Image.new("RGB", (W, H), "white")
        self.d = ImageDraw.Draw(self.im)
        self.number = number
        self.y = 96
        self.d.rectangle((0, 0, W, 22), fill=BLUE)
        self.d.text((100, 48), "01 版本白底窗帘成品图生成 SOP", font=f(23), fill=MUTED)
        self.d.text((100, self.y), title, font=f(55, True), fill=INK)
        self.y += 80
        if subtitle:
            self.d.text((100, self.y), subtitle, font=f(27), fill=MUTED)
            self.y += 55

    def section(self, text):
        self.y += 12
        self.d.text((100, self.y), text, font=f(35, True), fill=BLUE)
        self.y += 54

    def paragraph(self, text, size=28, color=INK, indent=100, width=1500, gap=12):
        font = f(size)
        line_h = int(size * 1.48)
        for line in wrap(self.d, text, font, width):
            self.d.text((indent, self.y), line, font=font, fill=color)
            self.y += line_h
        self.y += gap

    def bullets(self, items, size=27, indent=125, width=1450, gap=8):
        font = f(size)
        line_h = int(size * 1.42)
        for text in items:
            lines = wrap(self.d, text, font, width - 44)
            self.d.ellipse((indent, self.y + 15, indent + 12, self.y + 27), fill=BLUE)
            for i, line in enumerate(lines):
                self.d.text((indent + 34, self.y), line, font=font, fill=INK)
                self.y += line_h
            self.y += gap

    def numbered(self, items, size=27, indent=110, width=1480, gap=10):
        font = f(size)
        line_h = int(size * 1.42)
        for n, text in enumerate(items, 1):
            tag = f"{n}."
            self.d.text((indent, self.y), tag, font=font, fill=BLUE)
            lines = wrap(self.d, text, font, width - 58)
            for line in lines:
                self.d.text((indent + 54, self.y), line, font=font, fill=INK)
                self.y += line_h
            self.y += gap

    def callout(self, label, text, kind="info"):
        color = {"info": LIGHT, "warn": WARN, "risk": RISK, "success": "#E9F7EE"}[kind]
        label_color = {"info": BLUE, "warn": "#7A5A00", "risk": "#9B1C1C", "success": GREEN}[kind]
        font = f(26)
        text_lines = wrap(self.d, text, font, 1240)
        h = 34 + len(text_lines) * 40 + 20
        self.d.rounded_rectangle((100, self.y, 1600, self.y + h), radius=16, fill=color, outline=BORDER, width=2)
        self.d.text((128, self.y + 22), label, font=f(27, True), fill=label_color)
        tx = 355
        ty = self.y + 22
        for line in text_lines:
            self.d.text((tx, ty), line, font=font, fill=INK)
            ty += 40
        self.y += h + 20

    def table(self, headers, rows, widths, size=23):
        x = 100
        top = self.y
        font = f(size)
        bold = f(size, True)
        cols = [int(v) for v in widths]
        header_h = 54
        self.d.rounded_rectangle((x, top, x + sum(cols), top + header_h), radius=10, fill=LIGHT, outline=BORDER, width=2)
        cur = x
        for value, cw in zip(headers, cols):
            for i, line in enumerate(wrap(self.d, value, bold, cw - 22)):
                self.d.text((cur + 12, top + 12 + i * 27), line, font=bold, fill=INK)
            cur += cw
        y = top + header_h
        for row in rows:
            wrapped = [wrap(self.d, str(v), font, cw - 22) for v, cw in zip(row, cols)]
            rh = max(48, max(len(lines) for lines in wrapped) * 31 + 20)
            self.d.rectangle((x, y, x + sum(cols), y + rh), fill="white", outline=BORDER, width=2)
            cur = x
            for lines, cw in zip(wrapped, cols):
                for i, line in enumerate(lines):
                    self.d.text((cur + 12, y + 10 + i * 31), line, font=font, fill=INK)
                self.d.line((cur + cw, y, cur + cw, y + rh), fill=BORDER, width=2)
                cur += cw
            y += rh
        self.y = y + 18

    def finish(self):
        self.d.line((100, H - 88, 1600, H - 88), fill=BORDER, width=2)
        self.d.text((100, H - 66), "内部操作手册 · 2026-08-19", font=f(20), fill=MUTED)
        self.d.text((1510, H - 66), str(self.number), font=f(20), fill=MUTED)
        return self.im


def add_page(pages, number, title, subtitle, content):
    p = Page(number, title, subtitle)
    content(p)
    pages.append(p.finish())


def build_pages():
    pages = []
    add_page(pages, 1, "白底窗帘成品图生成 SOP", "适用于 01 版本：实拍面料图、CSV 花位数据与完整花位图", lambda p: (
        p.callout("核心目标", "在纯白背景上输出正视、完整、中间对开的双片窗帘。Image 1 只提供结构；Image 2 只提供面料颜色与材质；有花位时 Image 3 只提供完整花位的二维结构。", "info"),
        p.section("固定成品参数"),
        p.table(["项目", "固定值 / 规则"], [
            ("成品闭合尺寸", "整套宽 330 cm × 高 270 cm；左右片成品宽各 165 cm。"),
            ("褶皱", "2 倍褶皱。每片展开布为 330 cm × 270 cm；两片展开总布宽 660 cm。"),
            ("中间开口", "约 30 cm；开口区域没有面料，不额外增加成品额定宽度。"),
            ("左右裁片", "来自同一卷连续面料的相邻连续位置；不可复制、镜像、轴对称或只改亮度。"),
            ("纯白背景", "只属于背景，不得漂白、灰化或改变 Image 2 的任何固有面料颜色。"),
        ], [350, 1150], 24),
        p.section("适用资料"),
        p.bullets([
            "01 版本白底图生成页面，常用地址：http://127.0.0.1:3456。",
            "32-10：每张实拍面料图覆盖真实 10 cm × 10 cm；32-15：覆盖真实 15 cm × 15 cm。",
            "32 文件夹中的 CSV 管理 SKU、花位纵横尺寸和成分；居莱/花位图保存有花位 SKU 的完整花回图。",
        ])
    ))

    add_page(pages, 2, "一、输入职责与流程分流", "先看 CSV，再决定是否启用完整花位图", lambda p: (
        p.section("三张图和 CSV 的职责"),
        p.table(["输入", "只负责什么", "不得负责什么"], [
            ("Image 1\n窗帘结构模板", "帘头、孔位/捏褶、轮廓、两片布局、开口、褶峰褶谷、垂坠与宏观明暗。", "面料颜色、图案、织纹、材质。"),
            ("Image 2\n实拍面料图", "固有颜色、色彩比例、经纬、纱线、颗粒、毛羽、孔隙、光泽与质感。", "帘头、开口、褶皱数量和窗帘几何。"),
            ("Image 3\n完整花位图", "一个完整花位的二维结构、方向、元素位置和完整边界。仅用于有花位。", "颜色、亮度和材质；这些一律以 Image 2 为准。"),
            ("CSV", "SKU、是否有花位、花位纵横厘米尺寸、成分等辅助信息。", "不能替代 Image 2 的视觉证据。"),
        ], [300, 690, 610], 21),
        p.section("分流规则"),
        p.table(["CSV 花位字段", "执行流程", "提交资料"], [
            ("纵向长 = / 且横向长 = /", "无花位 / 素面微纹理", "Image 1 + Image 2；不得凭空生成条纹、方格或花位。"),
            ("两个字段均为有效正数", "有花位 / 三图流程", "Image 1 + Image 2 + 对应 SKU 的 Image 3；按 CSV 花回厘米计算循环。"),
            ("仅一个字段有值或不是正数", "阻止生成", "先修正 CSV；不得靠提示词或图像猜测花位尺寸。"),
        ], [520, 430, 650], 21),
        p.callout("关键原则", "10/15 cm 实拍范围只锁定材质密度，不等于花回尺寸；有花位时，真正的完整循环尺寸只读取 CSV 的纵向长和横向长。", "warn")
    ))

    add_page(pages, 3, "二、无花位 / 素面微纹理 SOP", "适用条件：CSV 的花位纵向长和横向长都为“/”", lambda p: (
        p.callout("目标", "素面不代表无细节。Image 2 的色纱、经纬、孔隙、毛羽、暗缝和微观反差必须保留；只能允许小细节在远景自然融合，不能被磨成平滑素面。", "info"),
        p.section("操作步骤"),
        p.numbered([
            "进入“白底图生成”，填写 API Key，并选择圆孔帘或韩式双捏褶帘。",
            "上传或选择 Image 2；确认 SKU、文件名以及 _detail_10 / _detail_15 对应的真实实拍范围。",
            "确认预检为“无花位”。无花位不匹配 Image 3；若 CSV 不为 / + /，停止并先核对 CSV。",
            "选择画质与分辨率后提交。默认提示词会将 Image 1 结构与 Image 2 材质彻底分离。",
            "完成后先看综合色相和织物质感；不合格时重新生成或补充材质要求，不要把样本手工平铺成规律花回。",
        ]),
        p.section("验收清单"),
        p.bullets([
            "全景颜色仍读取为 Image 2 的综合色相，不能被纯白背景洗成统一灰、米白、卡其或冷白。",
            "经纬、粗细差、毛羽、孔隙、暗缝和微高光可自然融合，但不能变成平滑纸面、塑料或通用噪点。",
            "不得新增条纹、方格、棋盘格、方框、网眼或独立大花位。",
            "褶谷阴影不能替代固有深色；褶峰高光不能漂白固有色纱。",
            "左右片同方向、同密度、同色系，但局部纹理不完全镜像。",
        ]),
        p.callout("尺寸提示", "10 cm 样本在 270 cm 高度中对应 27 个材质密度跨度；15 cm 对应 18 个跨度。该换算只用于保持纱线和纹理频率，不能据此虚构花回。", "warn")
    ))

    add_page(pages, 4, "三、有花位 / 完整花回 SOP", "适用条件：CSV 花位纵向长和横向长均为有效正数", lambda p: (
        p.callout("物理尺度优先", "Image 3 回答“一个花位长什么样”；CSV 回答“这个花位实际多大”；Image 2 回答“它用什么颜色和材质织成”。三者必须同时成立。", "warn"),
        p.section("操作步骤"),
        p.numbered([
            "确认 CSV 两个花位厘米字段均为正数；确认居莱/花位图中存在与 SKU 精确匹配的完整花位图。",
            "上传 Image 2。预检必须显示“有花位 X × Y cm · Image 3 已匹配”；缺图或字段异常时不得提交。",
            "固定展开尺寸：成品高 270 cm；每片展开宽 330 cm；两片总展开宽 660 cm。",
            "最终提示词必须写明：纵向花回数 = 270 ÷ 纵向花位 cm；每片横向花回数 = 330 ÷ 横向花位 cm；总横向花回数 = 660 ÷ 横向花位 cm。",
            "先生成连续平展花位面料，再裁切左右两片，最后包覆到 Image 1 褶皱；不得逐褶重置图案。",
        ]),
        p.section("有花位验收"),
        p.bullets([
            "花回数量符合 CSV 换算；边缘允许有裁切的部分花回，但不得少循环或把花回放大。",
            "花位穿过褶峰、侧面、褶谷时连续压缩和遮挡，不得在褶谷切断、镜像或重新开始。",
            "Image 3 的结构和 Image 2 的颜色、织纹、色纱、质感同时存在。",
            "左右片来自同一卷连续面料的相邻位置；花位相位合理但不是完全相同。",
        ])
    ))

    add_page(pages, 5, "四、尺度、密度与提示词检查", "花回尺度、实拍材质尺度与褶皱尺度必须分开", lambda p: (
        p.section("三种尺度的正确用途"),
        p.table(["尺度来源", "正确用途", "禁止错误"], [
            ("CSV 花位纵横厘米", "计算完整花回在 270 / 330 / 660 cm 展开布上的数量。", "把实拍图可见的局部数量当作完整花回数量。"),
            ("Image 3 完整花位图", "确定一个花回的轮廓、内部关系、方向和边界。", "用 Image 3 的截图光线、色偏或平滑纹理覆盖 Image 2。"),
            ("Image 2 10/15 cm 范围", "确定色纱比例、织造颗粒、纱线粗细、孔隙和真实材质频率。", "拉伸成花回、当矩形贴图平铺、或放大微观网眼。"),
        ], [480, 600, 520], 20),
        p.section("提示词硬规则"),
        p.bullets([
            "Image 1 只控制帘头、孔位/捏褶、完整轮廓、两片布局、留白、褶皱数量、褶峰褶谷和阴影明暗。",
            "Image 2 是唯一面料来源，控制颜色、深浅比例、织纹、颗粒、肌理、材料细节和真实色差。",
            "有花位时 Image 3 只控制完整花位结构与方向；CSV 的花位尺寸是绝对物理尺度。",
            "纯白只用于背景。不得偏色、提亮、压暗、去色，也不得将面料简化成通用条纹、网格或噪点。",
            "纹理低于全景像素时应自然融合为材料颗粒和综合色差，不能为了清晰被描粗成大网眼或大图案。",
        ]),
        p.section("快速审图顺序"),
        p.numbered([
            "先数显著花回或横带：明显少于 270 ÷ 纵向周期，说明图案被错误放大。",
            "再看综合色相：成品必须保留 Image 2 的蓝绿、砖红、黄赭、棕灰等固有色，不得统一泛灰。",
            "最后看材质：纱线、孔隙、暗缝和粗糙度在远景可融合，但综合色差和织物厚薄感不能消失。",
        ])
    ))

    add_page(pages, 6, "五、批量生成与导出保存", "手动上传和数据库批量共用同一套预检逻辑", lambda p: (
        p.section("手动上传批量"),
        p.numbered([
            "选择帘头样式，选择多个 Image 2 文件或一个面料文件夹。",
            "确认预检：无花位显示“/ + /”；有花位显示纵向 × 横向厘米及 Image 3 已匹配。",
            "字段异常、完整花位图缺失或 SKU 未匹配时，先处理数据，不要直接生成。",
            "选择画质和分辨率，点击“开始生成”；完成任务进入“已完成”。",
        ]),
        p.section("数据库批量"),
        p.numbered([
            "切换“数据库批量”，选择任务方式与帘头样式。",
            "先选择保存目录，再点击“读取数据库”；只切换区域不会自动读取。",
            "确认任务数、输出格式和并发数；上游限流或失败时先降低并发。",
            "“继续未完成”会跳过已有进度；“重新生成全部”可能覆盖同名结果，需谨慎。",
        ]),
        p.section("导出 PNG / WebP"),
        p.callout("保存不了怎么办", "点击“按批次导出 PNG/WebP”后，先在“选择批次”里勾选任务，再点击底部“选择文件夹并保存 (N)”。最后必须在系统弹出的文件夹选择器中选定实际目录并确认；未完成系统选择不会写出文件。", "warn"),
        p.bullets([
            "若提示输出文件名重复，先更换目标目录或处理同名文件，避免误覆盖。",
            "圆孔帘默认命名：{sku}_grommet.webp；韩式双捏褶帘默认命名：{sku}_double-pinch.webp。",
        ])
    ))

    add_page(pages, 7, "六、常见问题与最终检查", "发生问题时先核对输入职责、CSV 和物理尺度，再调提示词", lambda p: (
        p.section("快速排障"),
        p.table(["现象", "优先检查", "处理"], [
            ("颜色发灰、发米、被洗淡", "Image 2 是否正确；纯白背景是否影响面料。", "强调保留固有综合色相与色彩面积比例；确保使用最新默认提示词。"),
            ("纹理像被磨平", "实拍图清晰度和材质描述是否完整。", "强调经纬、毛羽、孔隙、暗缝和粗细差自然融合，不能平滑。"),
            ("花回数量不对", "CSV 的纵横花位厘米、Image 3 SKU 是否精确匹配。", "按 270/V、330/H、660/H 复核；不使用样本尺寸代替花回。"),
            ("无花位却生成条纹/方格", "CSV 是否为 / + /，是否误发 Image 3。", "只走两图流程，明确不新增中低频装饰图案。"),
            ("左右片完全相同", "生成是否出现镜像或复制。", "重生图，强调同卷相邻连续位置、不同裁切起点。"),
            ("预检/专项分析报错", "SKU、_detail_10/_detail_15、CSV 或 Image 3 缺失。", "按报错字段修正后再重新分析；不要绕过异常直接提交。"),
        ], [340, 530, 730], 19),
        p.section("提交前最终检查"),
        p.bullets([
            "SKU、Image 2、实拍范围（10/15 cm）和帘头样式均正确。",
            "无花位：CSV 两字段都是 /，只使用 Image 1 + Image 2。",
            "有花位：CSV 两字段均为正数，Image 3 已按 SKU 精确匹配。",
            "Image 1 只结构；Image 2 只颜色与材质；Image 3 只花位结构。",
            "已核对 270 cm 高度、330 cm 单片展开宽、660 cm 总展开宽的花回/密度换算。",
            "导出前确认批次、格式、目标文件夹和重名风险。",
        ]),
        p.callout("完成标准", "输出必须为正视、完整、中间对开的双片成品窗帘，纯白无纹理背景；不出现房间、窗户、轨道、挂钩、文字、水印、尺寸线、投影或商品卡片效果。", "success")
    ))
    return pages


def build_docx(pages):
    PAGE_DIR.mkdir(parents=True, exist_ok=True)
    for old in PAGE_DIR.glob("page-*.png"):
        old.unlink()
    paths = []
    for index, page in enumerate(pages, 1):
        path = PAGE_DIR / f"page-{index}.png"
        page.save(path, quality=95)
        paths.append(path)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.header_distance = Inches(0)
    section.footer_distance = Inches(0)
    for i, path in enumerate(paths):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Inches(0)
        paragraph.paragraph_format.space_before = Inches(0)
        paragraph.add_run().add_picture(str(path), width=Inches(7.5))
        if i < len(paths) - 1:
            doc.add_page_break()
    doc.core_properties.title = "01版本白底窗帘生图SOP"
    doc.core_properties.subject = "白底窗帘成品图生成操作标准流程"
    doc.core_properties.author = "Codex"
    doc.save(OUT)


if __name__ == "__main__":
    build_docx(build_pages())
    print(OUT)
